from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"F:\tool share\产品说明")
SCREENSHOT_DIR = ROOT / "screenshots"
OUTPUT = ROOT / "企具共享-产品使用说明.docx"

NAVY = RGBColor(24, 49, 83)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 101, 115)
LIGHT_FILL = "E8EEF5"
BORDER = "D9E2EC"


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def style_paragraph(paragraph, before=0, after=6, line=1.25, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "企具共享 Tool Share"
    style_paragraph(header, after=0)
    set_run_font(header.runs[0], size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = "产品使用说明"
    set_run_font(footer.runs[0], size=9, color=MUTED)


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("Tool Share 使用指南")
    set_run_font(run, size=10.5, bold=True, color=BLUE)
    style_paragraph(kicker, after=18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("企具共享产品使用说明")
    set_run_font(run, size=28, bold=True, color=NAVY)
    style_paragraph(title, after=8)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("用简单步骤说明如何找工具、提交内容、审核和管理平台")
    set_run_font(run, size=13, color=MUTED)
    style_paragraph(subtitle, after=30)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"适用对象：普通用户、审核员、系统管理员  |  生成日期：{date.today().isoformat()}")
    set_run_font(run, size=10.5, color=MUTED)

    doc.add_page_break()


def add_intro(doc):
    doc.add_heading("一、先认识这个系统", level=1)
    p = doc.add_paragraph()
    p.add_run("企具共享").bold = True
    p.add_run("是企业内部的工具共享平台。大家可以在这里查找常用工具、收藏好用工具、提交新工具，也可以把多个工具整理成一套工作流。管理员还可以维护账号、标签、审核、导入导出和日志。")
    style_paragraph(p)
    for run in p.runs:
        set_run_font(run)

    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Inches(1.5), Inches(2.3), Inches(2.7)]
    headers = ["身份", "适合谁使用", "主要能做什么"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_shading(cell, LIGHT_FILL)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(text)
        set_run_font(run, bold=True, color=NAVY)
    rows = [
        ("普通用户", "所有使用工具的同事", "浏览、搜索、收藏、评论、提交工具和工作流"),
        ("审核员", "负责内容质量的人", "审核工具和工作流，通过或驳回提交"),
        ("系统管理员", "平台维护人员", "管理账号、标签、工具、工作流、导入导出和日志"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            run = cells[i].paragraphs[0].add_run(text)
            set_run_font(run)


def add_image(doc, filename, caption):
    image_path = SCREENSHOT_DIR / filename
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(6.3))
    style_paragraph(p, before=4, after=3)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_run_font(run, size=9.5, color=MUTED, italic=True)
    style_paragraph(cap, after=10)


def add_steps(doc, title, steps):
    doc.add_heading(title, level=2)
    for step in steps:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(step)
        set_run_font(run)
        style_paragraph(p, after=4)


def add_bullets(doc, title, bullets):
    doc.add_heading(title, level=3)
    for item in bullets:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)
        style_paragraph(p, after=4)


def add_user_section(doc):
    doc.add_heading("二、普通用户怎么用", level=1)
    add_image(doc, "02-tools.png", "图 1：工具目录是用户查找和进入工具的主要入口。")

    add_steps(doc, "查找工具", [
        "进入“工具目录”。",
        "在搜索框输入工具名称、用途或关键词。",
        "如果工具很多，可以按标签筛选，也可以按时间或热度排序。",
        "点击工具名称进入详情页，查看链接、说明、使用方法和评论。",
    ])
    add_bullets(doc, "常用操作", [
        "觉得工具有用，可以在详情页点赞或收藏。",
        "收藏后的工具会集中出现在“我的收藏”。",
        "有使用经验、注意事项或推荐理由，可以在评论区补充。",
    ])

    add_image(doc, "03-submit.png", "图 2：提交工具时，按页面字段填写名称、摘要、描述、标签、链接和使用方法。")
    add_steps(doc, "提交新工具", [
        "进入“提交工具”。",
        "填写工具名称、简短摘要、详细描述、标签、工具链接和使用方法。",
        "确认链接以 http 或 https 开头，名称不要和已有工具重复。",
        "点击“提交工具”。如果开启审核，提交后会先进入待审核状态。",
    ])

    add_steps(doc, "提交工作流", [
        "进入“提交工具”，切换到“提交工作流”。",
        "填写工作流名称、适用场景、描述和步骤说明。",
        "选择至少一个关联工具。",
        "提交后等待审核员或管理员审核。",
    ])
    add_image(doc, "04-workflows.png", "图 3：工作流用于把多个工具串成一套可复用的方法。")


def add_reviewer_section(doc):
    doc.add_heading("三、审核员怎么用", level=1)
    p = doc.add_paragraph("审核员主要负责判断提交内容是否真实、清楚、可用。审核时不要只看标题，要重点看描述、使用方法、标签和关联工具是否合理。")
    for run in p.runs:
        set_run_font(run)
    style_paragraph(p)

    add_image(doc, "05-reviews.png", "图 4：审核中心可以处理工具审核和工作流审核。")
    add_steps(doc, "审核工具或工作流", [
        "进入“审核中心”。",
        "选择工具审核或工作流审核。",
        "筛选“待审核”内容。",
        "逐条查看提交信息，确认内容是否完整、链接是否有效、步骤是否清楚。",
        "符合要求就点击“通过”；不符合要求就点击“驳回”，并写清楚原因。",
    ])
    add_bullets(doc, "驳回建议", [
        "原因要具体，例如“使用方法不清楚”“链接打不开”“标签不合适”。",
        "如果内容有价值但表达不完整，建议写出修改方向，方便提交人一次改对。",
    ])


def add_admin_section(doc):
    doc.add_heading("四、系统管理员怎么用", level=1)
    p = doc.add_paragraph("系统管理员拥有全量管理权限，适合负责平台初始化、账号开通、基础数据维护、审核策略和日志追踪。")
    for run in p.runs:
        set_run_font(run)
    style_paragraph(p)

    add_image(doc, "06-users.png", "图 5：账号权限页面用于新增账号、编辑角色、停用账号和重置密码。")
    add_steps(doc, "管理账号", [
        "进入“账号权限”。",
        "点击“新增账号”，填写用户名、真实姓名、初始密码、状态和角色。",
        "如果给同事分配 ADMIN 角色，该账号会自动拥有管理员权限。",
        "已有账号可以编辑、停用或重置密码。",
    ])

    add_steps(doc, "维护工具、标签和工作流", [
        "进入“工具管理”，可以直接新增、编辑或删除工具。",
        "进入“标签管理”，可以新增、编辑或删除标签。",
        "在工具管理页进入“管理工作流”，可以维护工作流并设置精选内容。",
    ])

    add_image(doc, "07-import-export.png", "图 6：导入导出页面适合批量导入工具或导出平台数据。")
    add_steps(doc, "批量导入和导出", [
        "进入“导入导出”。",
        "导入前先下载模板，按模板整理 CSV 数据。",
        "上传后先做预校验，确认无误后再正式导入。",
        "需要盘点或归档时，可以选择字段并导出工具数据。",
    ])

    add_image(doc, "08-logs.png", "图 7：日志审计页面用于查看登录、审核、导入导出等关键操作记录。")
    add_steps(doc, "查看日志", [
        "进入“日志审计”。",
        "查看审计日志或操作日志。",
        "按动作、模块、成功状态或操作者筛选。",
        "排查问题时，优先查看失败记录和最近操作记录。",
    ])


def add_faq(doc):
    doc.add_heading("五、常见问题", level=1)
    faqs = [
        ("提交工具后多久能发布？", "如果工具审核关闭，会自动发布；如果审核开启，需要审核员或管理员通过后才会发布。"),
        ("工具被驳回怎么办？", "查看驳回原因，按说明修改后重新提交。"),
        ("忘记密码怎么办？", "联系管理员在“账号权限”里重置密码。"),
        ("工作流一定要关联工具吗？", "是。工作流至少要关联一个已有工具。"),
        ("标签可以随便建吗？", "不建议。标签应由管理员统一维护，避免重复和含义混乱。"),
    ]
    for question, answer in faqs:
        p = doc.add_paragraph()
        q = p.add_run(f"Q：{question}\n")
        set_run_font(q, bold=True, color=NAVY)
        a = p.add_run(f"A：{answer}")
        set_run_font(a)
        style_paragraph(p, after=8)


def main():
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_intro(doc)
    add_user_section(doc)
    add_reviewer_section(doc)
    add_admin_section(doc)
    add_faq(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
